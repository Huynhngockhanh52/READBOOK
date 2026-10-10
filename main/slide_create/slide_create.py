#!/usr/bin/env python3
"""Create PowerPoint files from validated slide-section Markdown."""

from __future__ import annotations

import argparse
import re
import shutil
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence
from urllib.parse import unquote


yaml_err: Exception | None = None
try:
    import yaml
except ImportError as exc:
    yaml = None
    yaml_err = exc

pptx_err: Exception | None = None
try:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.oxml.xmlchemy import OxmlElement
except ImportError as exc:
    Presentation = None
    PP_PLACEHOLDER = None
    OxmlElement = None
    pptx_err = exc

try:
    from sum2slide import SlideData, parseSlides, validateSlides
except ImportError:
    from .sum2slide import SlideData, parseSlides, validateSlides


IMAGE_RE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
LIST_RE = re.compile(r"^(\s*)([-+*]|\d+[.)])\s+(.+)$")
HEAD_RE = re.compile(r"^#{1,6}\s+(.+)$")
TABLE_SEP = re.compile(r"^\s*\|?(?:\s*:?-{3,}:?\s*\|)+\s*$")
INLINE_RE = re.compile(
    r"(\*\*\*.+?\*\*\*|___[^_]+?___|\*\*.+?\*\*|__[^_]+?__|"
    r"`[^`]+?`|\*[^*]+?\*|_[^_]+?_|\[[^\]]+\]\([^)]+\))"
)
IMG_EXTS = {".bmp", ".gif", ".jpeg", ".jpg", ".png", ".tif", ".tiff", ".wmf"}


class CreateError(RuntimeError):
    """Raised when a PowerPoint cannot be created safely."""


@dataclass(frozen=True)
class TextBlock:
    text: str
    level: int = 0
    bullet: bool = False
    number: int | None = None


@dataclass(frozen=True)
class InlinePart:
    text: str
    bold: bool = False
    italic: bool = False


def setUtf8() -> None:
    for stream in (sys.stdout, sys.stderr):
        reconfig = getattr(stream, "reconfigure", None)
        if reconfig is not None:
            reconfig(encoding="utf-8", errors="replace")


def getRoot() -> Path:
    return Path(__file__).resolve().parents[2]


def parseArgs(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Tạo PowerPoint từ Markdown slide đã chuẩn hóa.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument(
        "-i",
        "--input",
        dest="input_path",
        type=Path,
        required=True,
        help="File Markdown chứa nội dung slide.",
    )
    parser.add_argument(
        "-o",
        "--output",
        dest="output_path",
        type=Path,
        help="File .pptx hoặc thư mục output tùy chọn.",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Ghi đè PowerPoint đã tồn tại.",
    )
    return parser.parse_args(argv)


def needDeps() -> None:
    missing: list[str] = []
    if yaml_err is not None or yaml is None:
        missing.append("PyYAML")
    if (
        pptx_err is not None
        or Presentation is None
        or PP_PLACEHOLDER is None
        or OxmlElement is None
    ):
        missing.append("python-pptx")
    if missing:
        packages = " ".join(missing)
        cause = pptx_err or yaml_err
        raise CreateError(f"Thiếu dependency. Hãy chạy: pip install {packages}") from cause


def getTemplate() -> Path:
    path = getRoot() / "sample" / "base_book.pptx"
    if not path.is_file():
        raise CreateError(f"Không tìm thấy PowerPoint template: {path}")
    return path


def loadLayouts() -> dict[str, tuple[str, ...]]:
    path = Path(__file__).resolve().parent / "conf" / "base_book.yaml"
    if not path.is_file():
        raise CreateError(f"Không tìm thấy layout config: {path}")
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, yaml.YAMLError) as exc:
        raise CreateError(f"Không thể đọc YAML {path}: {exc}") from exc
    if not isinstance(data, dict) or not isinstance(data.get("layouts"), dict):
        raise CreateError("base_book.yaml thiếu mapping layouts.")

    layouts: dict[str, tuple[str, ...]] = {}
    for raw_name, raw_fields in data["layouts"].items():
        if not isinstance(raw_name, str) or not raw_name.strip():
            raise CreateError(f"Tên layout không hợp lệ: {raw_name!r}")
        name = raw_name.strip()
        if not isinstance(raw_fields, list):
            raise CreateError(f"layouts[{name!r}] phải là list placeholder.")
        fields: list[str] = []
        for raw_field in raw_fields:
            if not isinstance(raw_field, str) or not raw_field.strip():
                raise CreateError(f"Placeholder không hợp lệ trong layout {name!r}.")
            field = raw_field.strip()
            if field in fields:
                raise CreateError(f"Placeholder {field!r} bị lặp trong layout {name!r}.")
            fields.append(field)
        layouts[name] = tuple(fields)
    if not layouts:
        raise CreateError("Không có layout nào trong base_book.yaml.")
    return layouts


def getInput(path: Path) -> Path:
    in_path = path.expanduser().resolve()
    if not in_path.is_file():
        raise CreateError(f"Không tìm thấy Markdown input: {in_path}")
    if in_path.suffix.casefold() != ".md":
        raise CreateError(f"Input phải là file .md: {in_path}")
    return in_path


def pptxName(path: Path) -> str:
    stem = path.stem
    if stem.casefold().endswith("_slide"):
        stem = stem[:-6]
    return f"{stem}.pptx"


def getOutput(in_path: Path, out_arg: Path | None) -> Path:
    if out_arg is None:
        out_path = in_path.parent.parent / "pptx" / pptxName(in_path)
    else:
        target = out_arg.expanduser().resolve()
        if target.suffix.casefold() == ".pptx" and not target.is_dir():
            out_path = target
        else:
            out_path = target / pptxName(in_path)
    if out_path.suffix.casefold() != ".pptx":
        raise CreateError(f"Output phải có đuôi .pptx: {out_path}")
    if out_path.exists() and out_path.is_dir():
        raise CreateError(f"Output đang là thư mục: {out_path}")
    return out_path


def loadSlides(
    in_path: Path,
    layouts: dict[str, tuple[str, ...]],
) -> list[SlideData]:
    try:
        markdown = in_path.read_text(encoding="utf-8")
    except (OSError, UnicodeError) as exc:
        raise CreateError(f"Không thể đọc Markdown {in_path}: {exc}") from exc
    try:
        slides = parseSlides(markdown)
        validateSlides(slides, layouts)
    except Exception as exc:
        if isinstance(exc, CreateError):
            raise
        raise CreateError(f"Markdown slide không hợp lệ: {exc}") from exc
    return slides


def layoutMap(prs: Any) -> dict[str, Any]:
    layouts: dict[str, Any] = {}
    duplicates: set[str] = set()
    for layout in prs.slide_layouts:
        name = str(layout.name or "")
        if name in layouts:
            duplicates.add(name)
        else:
            layouts[name] = layout
    if duplicates:
        names = ", ".join(sorted(duplicates))
        raise CreateError(f"PowerPoint template có layout trùng tên: {names}")
    return layouts


def checkLayouts(
    slides: list[SlideData],
    ppt_layouts: dict[str, Any],
) -> None:
    used = {slide.layout for slide in slides}
    missing = sorted(name for name in used if name not in ppt_layouts)
    if missing:
        names = ", ".join(missing)
        raise CreateError(f"Template PowerPoint thiếu layout: {names}")


def clearSlides(prs: Any) -> None:
    slide_ids = list(prs.slides._sldIdLst)
    for slide_id in slide_ids:
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)


def syncShapeNames(slide: Any, layout: Any, slide_num: int) -> None:
    """Copy placeholder names from the layout by stable placeholder idx."""
    layout_by_idx: dict[int, str] = {}
    name_to_idx: dict[str, int] = {}

    for layout_shape in layout.shapes:
        if not layout_shape.is_placeholder:
            continue
        idx = int(layout_shape.placeholder_format.idx)
        name = str(layout_shape.name or "").strip()
        if not name:
            raise CreateError(
                f"Slide {slide_num}, layout {layout.name!r}: "
                f"placeholder idx={idx} không có tên."
            )
        if idx in layout_by_idx:
            raise CreateError(
                f"Slide {slide_num}, layout {layout.name!r}: "
                f"placeholder idx={idx} xuất hiện nhiều lần."
            )
        if name in name_to_idx:
            other_idx = name_to_idx[name]
            raise CreateError(
                f"Slide {slide_num}, layout {layout.name!r}: tên placeholder "
                f"{name!r} bị trùng tại idx={other_idx} và idx={idx}."
            )
        layout_by_idx[idx] = name
        name_to_idx[name] = idx

    slide_idxs: set[int] = set()
    for slide_shape in slide.shapes:
        if not slide_shape.is_placeholder:
            continue
        idx = int(slide_shape.placeholder_format.idx)
        if idx in slide_idxs:
            raise CreateError(
                f"Slide {slide_num}, layout {layout.name!r}: "
                f"placeholder idx={idx} bị trùng trên slide mới."
            )
        slide_idxs.add(idx)
        layout_name = layout_by_idx.get(idx)
        if layout_name is not None:
            slide_shape.name = layout_name


def shapeMap(slide: Any, slide_num: int, layout: str) -> dict[str, Any]:
    shapes: dict[str, Any] = {}
    duplicates: set[str] = set()
    for shape in slide.shapes:
        name = str(shape.name or "")
        if name in shapes:
            duplicates.add(name)
        else:
            shapes[name] = shape
    if duplicates:
        names = ", ".join(sorted(duplicates))
        raise CreateError(
            f"Slide {slide_num}, layout {layout!r}: shape trùng tên: {names}"
        )
    return shapes


def plainMd(text: str) -> str:
    value = IMAGE_RE.sub(lambda match: match.group(1), text)
    lines: list[str] = []
    for line in value.splitlines():
        clean_line = re.sub(r"^#{1,6}\s+", "", line)
        clean_line = "".join(part.text for part in inlineParts(clean_line))
        lines.append(clean_line)
    return "\n".join(lines).strip()


def textBlocks(text: str) -> list[TextBlock]:
    blocks: list[TextBlock] = []
    paragraph: list[str] = []
    in_code = False

    def flushPara() -> None:
        if paragraph:
            value = " ".join(item.strip() for item in paragraph if item.strip())
            if value:
                blocks.append(TextBlock(value))
            paragraph.clear()

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if re.match(r"^\s*(`{3,}|~{3,})", line):
            flushPara()
            in_code = not in_code
            continue
        if in_code:
            if line:
                blocks.append(TextBlock(line))
            continue
        if not line.strip():
            flushPara()
            continue
        if TABLE_SEP.match(line):
            continue
        table_cells = None
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            table_cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if table_cells is not None:
            flushPara()
            blocks.append(TextBlock(" • ".join(table_cells)))
            continue
        list_match = LIST_RE.match(line)
        if list_match:
            flushPara()
            indent = len(list_match.group(1).replace("\t", "    "))
            marker = list_match.group(2)
            level = min(indent // 2, 8)
            if marker[0].isdigit():
                number = int(re.match(r"\d+", marker).group())
                blocks.append(TextBlock(list_match.group(3), level, False, number))
            else:
                blocks.append(TextBlock(list_match.group(3), level, True))
            continue
        head_match = HEAD_RE.match(stripped)
        if head_match:
            flushPara()
            blocks.append(TextBlock(head_match.group(1)))
            continue
        paragraph.append(stripped)
    flushPara()
    return blocks


def inlineParts(text: str) -> list[InlinePart]:
    parts: list[InlinePart] = []
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            parts.append(InlinePart(text[cursor : match.start()]))
        token = match.group(0)
        if token.startswith("***") and token.endswith("***"):
            parts.append(InlinePart(token[3:-3], True, True))
        elif token.startswith("___") and token.endswith("___"):
            parts.append(InlinePart(token[3:-3], True, True))
        elif token.startswith("**") and token.endswith("**"):
            parts.append(InlinePart(token[2:-2], True, False))
        elif token.startswith("__") and token.endswith("__"):
            parts.append(InlinePart(token[2:-2], True, False))
        elif token.startswith("`") and token.endswith("`"):
            parts.append(InlinePart(token[1:-1]))
        elif token.startswith("["):
            link_match = LINK_RE.fullmatch(token)
            label = link_match.group(1) if link_match else token
            parts.append(InlinePart(label))
        else:
            parts.append(InlinePart(token[1:-1], False, True))
        cursor = match.end()
    if cursor < len(text):
        parts.append(InlinePart(text[cursor:]))
    return parts or [InlinePart(text)]


def clearBullets(para: Any) -> Any:
    p_pr = para._p.get_or_add_pPr()
    for tag in ("{http://schemas.openxmlformats.org/drawingml/2006/main}buNone",
                "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar",
                "{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum"):
        for child in list(p_pr.findall(tag)):
            p_pr.remove(child)
    return p_pr


def setBullet(para: Any, block: TextBlock) -> None:
    p_pr = clearBullets(para)
    if block.bullet:
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "•")
        p_pr.append(bullet)
    elif block.number is not None:
        bullet = OxmlElement("a:buAutoNum")
        bullet.set("type", "arabicPeriod")
        bullet.set("startAt", str(block.number))
        p_pr.append(bullet)
    else:
        p_pr.append(OxmlElement("a:buNone"))


def addInline(para: Any, text: str) -> None:
    for part in inlineParts(text):
        run = para.add_run()
        run.text = part.text
        if part.bold:
            run.font.bold = True
        if part.italic:
            run.font.italic = True


def fillText(shape: Any, content: str, label: str) -> None:
    if not getattr(shape, "has_text_frame", False):
        raise CreateError(f"Shape {label!r} không hỗ trợ text frame.")
    text_frame = shape.text_frame
    text_frame.clear()
    blocks = textBlocks(content)
    if not blocks:
        text_frame.paragraphs[0].text = ""
        return
    for index, block in enumerate(blocks):
        para = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        para.level = block.level
        setBullet(para, block)
        addInline(para, block.text)


def imageRef(content: str) -> str | None:
    match = IMAGE_RE.search(content)
    if match:
        return match.group(1).strip().strip("<>")
    lines = [line.strip() for line in content.splitlines() if line.strip()]
    if len(lines) != 1:
        return None
    value = lines[0].strip("`\"'").strip("<>")
    suffix = Path(unquote(value)).suffix.casefold()
    return value if suffix in IMG_EXTS else None


def resolveImage(raw_path: str, md_path: Path) -> Path:
    value = unquote(raw_path).strip()
    if re.match(r"^[a-z]+://", value, re.IGNORECASE):
        raise CreateError(f"Chỉ hỗ trợ hình local, không hỗ trợ URL: {value}")
    path = Path(value).expanduser()
    if path.is_absolute():
        candidates = [path]
    else:
        candidates = [
            md_path.parent / path,
            md_path.parent.parent / "md" / path,
            md_path.parent.parent / path,
        ]
    for candidate in candidates:
        resolved = candidate.resolve()
        if resolved.is_file():
            return resolved
    tried = ", ".join(str(path.resolve()) for path in candidates)
    raise CreateError(f"Không tìm thấy hình {raw_path!r}. Đã thử: {tried}")


def removeShape(shape: Any) -> None:
    element = shape._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def fillImage(slide: Any, shape: Any, path: Path, label: str) -> None:
    insert_picture = getattr(shape, "insert_picture", None)
    if callable(insert_picture):
        try:
            picture = insert_picture(str(path))
            picture.name = label
            return
        except Exception as exc:
            raise CreateError(f"Không thể chèn hình vào {label!r}: {exc}") from exc
    left, top = shape.left, shape.top
    width, height = shape.width, shape.height
    try:
        removeShape(shape)
        picture = slide.shapes.add_picture(str(path), left, top, width, height)
        picture.name = label
    except Exception as exc:
        raise CreateError(f"Không thể chèn hình vào shape {label!r}: {exc}") from exc


def notesFrame(slide: Any) -> Any:
    notes_slide = slide.notes_slide
    text_frame = getattr(notes_slide, "notes_text_frame", None)
    if text_frame is not None:
        return text_frame
    for shape in notes_slide.placeholders:
        if (
            getattr(shape, "has_text_frame", False)
            and shape.placeholder_format.type == PP_PLACEHOLDER.BODY
        ):
            return shape.text_frame
    raise CreateError("Không tìm thấy Notes text frame trong PowerPoint template.")


def setNotes(slide: Any, script: str) -> None:
    text_frame = notesFrame(slide)
    text_frame.clear()
    text_frame.paragraphs[0].text = plainMd(script)


def addSlide(
    prs: Any,
    slide_data: SlideData,
    layout: Any,
    md_path: Path,
) -> None:
    slide = prs.slides.add_slide(layout)
    syncShapeNames(slide, layout, slide_data.number)
    shapes = shapeMap(slide, slide_data.number, slide_data.layout)
    for field, content in slide_data.fields.items():
        shape = shapes.get(field)
        if shape is None:
            names = ", ".join(sorted(shapes))
            raise CreateError(
                f"Slide {slide_data.number}, layout {slide_data.layout!r}: "
                f"không tìm thấy shape {field!r}. Shapes hiện có: {names}"
            )
        raw_image = imageRef(content)
        if raw_image is not None:
            image_path = resolveImage(raw_image, md_path)
            fillImage(slide, shape, image_path, field)
        else:
            fillText(shape, content, field)
    setNotes(slide, slide_data.script)


def checkSaved(path: Path, slides: list[SlideData]) -> None:
    try:
        prs = Presentation(str(path))
    except Exception as exc:
        raise CreateError(f"PowerPoint output không mở lại được: {exc}") from exc
    if len(prs.slides) != len(slides):
        raise CreateError(
            f"PowerPoint output có {len(prs.slides)} slide; mong đợi {len(slides)}."
        )
    for index, slide in enumerate(prs.slides):
        try:
            notes = notesFrame(slide).text.strip()
        except Exception as exc:
            raise CreateError(f"Slide {index + 1}: không đọc lại được Speaker Notes: {exc}") from exc
        expected = plainMd(slides[index].script)
        if notes != expected:
            raise CreateError(f"Slide {index + 1}: Speaker Notes không khớp sau khi lưu.")


def buildPptx(
    template: Path,
    out_path: Path,
    md_path: Path,
    slides: list[SlideData],
) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=out_path.parent,
            prefix=f".{out_path.stem}.",
            suffix=".pptx",
            delete=False,
        ) as handle:
            temp_path = Path(handle.name)
        shutil.copy2(template, temp_path)
        prs = Presentation(str(temp_path))
        ppt_layouts = layoutMap(prs)
        checkLayouts(slides, ppt_layouts)
        clearSlides(prs)
        for slide_data in slides:
            addSlide(
                prs,
                slide_data,
                ppt_layouts[slide_data.layout],
                md_path,
            )
        prs.save(str(temp_path))
        checkSaved(temp_path, slides)
        temp_path.replace(out_path)
    except CreateError:
        if temp_path is not None and temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass
        raise
    except Exception as exc:
        if temp_path is not None and temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass
        raise CreateError(f"Không thể tạo PowerPoint: {exc}") from exc


def main(argv: Sequence[str] | None = None) -> int:
    setUtf8()
    args = parseArgs(argv)
    try:
        needDeps()
        in_path = getInput(args.input_path)
        out_path = getOutput(in_path, args.output_path)
        if out_path.exists() and not args.overwrite:
            print(f"[SKIP] Output đã tồn tại: {out_path}")
            return 0
        template = getTemplate()
        layouts = loadLayouts()
        slides = loadSlides(in_path, layouts)
        print(f"Input   : {in_path}")
        print(f"Template: {template}")
        print(f"Output  : {out_path}")
        print(f"Slides  : {len(slides)}")
        buildPptx(template, out_path, in_path, slides)
        print(f"[DONE] {out_path}")
        return 0
    except CreateError as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        print("[ERROR] Đã dừng bởi người dùng.", file=sys.stderr)
        return 130


if __name__ == "__main__":
    raise SystemExit(main())
